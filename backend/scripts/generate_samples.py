import os
from fpdf import FPDF
from docx import Document

os.makedirs("tests/assets", exist_ok=True)
os.makedirs("../frontend/assets/sample_resumes", exist_ok=True)


def generate_backend_pdf(output_path: str):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    # Name and Header
    pdf.set_font("Helvetica", "B", 18)
    pdf.cell(0, 10, "Alex Mercer", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 5, "alex.mercer@email.com | +1 (555) 234-5678 | San Francisco, CA", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 5, "linkedin.com/in/alexmercer-dev | github.com/alexmercer", align="C", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(5)

    # Summary
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "PROFESSIONAL SUMMARY", new_x="LMARGIN", new_y="NEXT")
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(2)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 5, 
        "Senior Backend Engineer with 5+ years of experience designing and scaling distributed systems, RESTful microservices, and asynchronous event pipelines. Expertise in Python, FastAPI, Django, PostgreSQL, and AWS cloud environments. Passionate about automated testing, performance tuning, and clean software architecture."
    )
    pdf.ln(4)

    # Technical Skills
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "TECHNICAL SKILLS", new_x="LMARGIN", new_y="NEXT")
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(2)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 5,
        "- Languages: Python, SQL, JavaScript, Bash\n"
        "- Frameworks & Tools: FastAPI, Django REST Framework, Celery, Redis, Docker, Git\n"
        "- Databases: PostgreSQL, MySQL, Redis, MongoDB\n"
        "- Cloud & DevOps: AWS (ECS, S3, RDS, Lambda), CI/CD (GitHub Actions), Linux, Nginx\n"
        "- Testing: Pytest, Unit Testing, TDD, Integration Testing"
    )
    pdf.ln(4)

    # Work Experience
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "WORK EXPERIENCE", new_x="LMARGIN", new_y="NEXT")
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(2)

    # Job 1
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Senior Backend Engineer | CloudScale Inc. (2022 - Present)", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 5,
        "- Architected and deployed 12+ high-throughput microservices using Python and FastAPI, handling 15M+ daily requests with 99.98% uptime.\n"
        "- Optimized PostgreSQL database queries and connection pooling, reducing p99 API latency by 42% across core billing endpoints.\n"
        "- Spearheaded migration of background tasks to Celery and Redis clusters, cutting asynchronous execution times by 35%.\n"
        "- Engineered automated CI/CD deployment pipelines on AWS ECS using GitHub Actions, enabling zero-downtime releases."
    )
    pdf.ln(3)

    # Job 2
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Backend Software Engineer | DataStream Labs (2019 - 2022)", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(0, 5,
        "- Developed secure RESTful APIs utilizing Django and PostgreSQL supporting 100k+ active SaaS subscribers.\n"
        "- Implemented Pytest test suites achieving 88% code coverage, substantially reducing regression defects in production.\n"
        "- Automated containerization workflows with Docker and configured monitoring with Prometheus and Grafana."
    )
    pdf.ln(4)

    # Education
    pdf.set_font("Helvetica", "B", 12)
    pdf.cell(0, 8, "EDUCATION", new_x="LMARGIN", new_y="NEXT")
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 5, "Bachelor of Science in Computer Science | University of California, Berkeley", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(0, 5, "Graduated 2019 | Relevant Coursework: Data Structures, Algorithms, Distributed Systems", new_x="LMARGIN", new_y="NEXT")

    pdf.output(output_path)
    print(f"Generated PDF: {output_path}")


def generate_frontend_docx(output_path: str):
    doc = Document()

    # Title / Header
    header = doc.add_paragraph()
    header.add_run("Elena Rostova\n").bold = True
    header.paragraph_format.space_after = 4
    contact = doc.add_paragraph("elena.rostova@example.com | (555) 789-0123 | New York, NY\nlinkedin.com/in/elena-rostova | github.com/erostova")
    contact.paragraph_format.space_after = 12

    # Summary
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(
        "Creative Full Stack and Frontend Developer with 4 years of experience delivering responsive, accessible, and high-performance web applications. Skilled in React.js, TypeScript, Next.js, Node.js, and modern styling solutions including Tailwind CSS. Dedicated to component modularity, web vitals optimization, and clean API design."
    )

    # Skills
    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph(
        "- Languages: TypeScript, JavaScript (ES6+), HTML5, CSS3, SQL, Python\n"
        "- Frontend: React.js, Next.js, Redux Toolkit, Tailwind CSS, Webpack, Vite\n"
        "- Backend & Storage: Node.js, Express.js, PostgreSQL, MongoDB, REST APIs\n"
        "- Testing & Tooling: Jest, React Testing Library, Git, Docker, Figma"
    )

    # Experience
    doc.add_heading("Work Experience", level=2)
    p1 = doc.add_paragraph()
    p1.add_run("Frontend Software Engineer | NovaTech Solutions (2021 - Present)\n").bold = True
    p1.add_run(
        "- Spearheaded frontend revamp of enterprise SaaS dashboard using React, TypeScript, and Tailwind CSS, improving Lighthouse score from 64 to 96.\n"
        "- Engineered reusable UI component library across 4 internal products, reducing developer delivery time by 30%.\n"
        "- Integrated complex REST APIs and GraphQL mutations, streamlining state management with Redux Toolkit.\n"
        "- Authored 120+ unit and integration tests with Jest, elevating test coverage to 85%."
    )

    p2 = doc.add_paragraph()
    p2.add_run("Junior Web Developer | PixelCraft Media (2019 - 2021)\n").bold = True
    p2.add_run(
        "- Built dynamic customer-facing web portals utilizing JavaScript, React, and CSS3, supporting 50k+ monthly visitors.\n"
        "- Collaborated with UX designers to translate wireframes into pixel-perfect responsive layouts.\n"
        "- Implemented form validations and API error boundary handling, reducing bounce rates by 18%."
    )

    # Education
    doc.add_heading("Education", level=2)
    p3 = doc.add_paragraph()
    p3.add_run("Bachelor of Science in Information Technology\n").bold = True
    p3.add_run("New York University (NYU) - 2019")

    doc.save(output_path)
    print(f"Generated DOCX: {output_path}")


if __name__ == "__main__":
    generate_backend_pdf("tests/assets/sample_backend_developer.pdf")
    generate_backend_pdf("../frontend/assets/sample_resumes/sample_backend_developer.pdf")
    generate_frontend_docx("tests/assets/sample_frontend_developer.docx")
    generate_frontend_docx("../frontend/assets/sample_resumes/sample_frontend_developer.docx")
